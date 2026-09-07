import csv
import io
from asyncio import to_thread
from typing import cast
from uuid import UUID, uuid4

import boto3
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.core.exceptions import NotFoundError
from app.modules.file_upload.models import UploadedFile
from app.modules.file_upload.repository import FileRepository
from app.modules.file_upload.schemas import PreviewResponse


class FileService:
    def __init__(self, db: AsyncSession, workspace_id: UUID) -> None:
        self.repo = FileRepository(db, workspace_id)
        self.workspace_id = workspace_id
        settings = get_settings()
        self.s3 = boto3.client(
            "s3",
            endpoint_url=settings.s3_endpoint_url,
            aws_access_key_id=settings.s3_access_key,
            aws_secret_access_key=settings.s3_secret_key,
            region_name=settings.s3_region,
        )
        self.bucket = settings.s3_bucket

    async def upload(self, filename: str, content_type: str, body: bytes) -> UploadedFile:
        key = f"{self.workspace_id}/{uuid4()}/{filename}"
        await to_thread(
            self.s3.put_object,
            Bucket=self.bucket,
            Key=key,
            Body=body,
            ContentType=content_type,
        )
        return await self.repo.add(
            UploadedFile(
                workspace_id=self.workspace_id,
                filename=filename,
                content_type=content_type,
                storage_key=key,
                size_bytes=len(body),
            )
        )

    def _download_sync(self, item: UploadedFile) -> bytes:
        body = self.s3.get_object(Bucket=self.bucket, Key=item.storage_key)["Body"].read()
        return cast(bytes, body)

    @staticmethod
    def parse_content(filename: str, content: bytes) -> tuple[list[str], list[dict[str, object]]]:
        if filename.lower().endswith(".xlsx"):
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            values = list(workbook.active.iter_rows(values_only=True))
            columns = [str(value or "") for value in values[0]] if values else []
            return columns, [dict(zip(columns, row, strict=False)) for row in values[1:]]
        encoding = "utf-8-sig"
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            text = content.decode("cp1251")
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        rows = [{str(key): value for key, value in row.items()} for row in reader]
        return list(reader.fieldnames or []), rows

    async def preview(self, file_id: UUID) -> PreviewResponse:
        item = await self.repo.get(file_id)
        if not item:
            raise NotFoundError("Файл не найден")
        content = await to_thread(self._download_sync, item)
        columns, rows = self.parse_content(item.filename, content)
        if item.filename.lower().endswith(".xlsx"):
            return PreviewResponse(columns=columns, rows=rows[:5])
        encoding = "utf-8-sig"
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            encoding = "cp1251"
            text = content.decode(encoding)
        sample = text[:8192]
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        rows = [dict(row) for _, row in zip(range(5), reader, strict=False)]
        return PreviewResponse(
            columns=list(reader.fieldnames or []), rows=rows, delimiter=delimiter, encoding=encoding
        )

    async def rows(self, file_id: UUID) -> tuple[UploadedFile, list[dict[str, object]]]:
        item = await self.repo.get(file_id)
        if not item:
            raise NotFoundError("Файл не найден")
        content = await to_thread(self._download_sync, item)
        _, rows = self.parse_content(item.filename, content)
        return item, rows
